"""Generates two small sample files for testing: a DBMS PDF and a SQL PPTX."""
from pathlib import Path

from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

out = Path(__file__).parent
PAGES = [
    ("Unit 1: ER Model", "An entity-relationship (ER) model describes data as entities, attributes and relationships. An entity is a real-world object such as a Student or a Course. Attributes describe an entity, for example Name or RollNo. A relationship connects entities, for example Student enrolls in Course. Cardinality defines how many instances can take part: one-to-one, one-to-many or many-to-many. A primary key uniquely identifies each entity instance."),
    ("Unit 2: Functional Dependencies, 1NF and 2NF", "A functional dependency X -> Y means that the value of X determines the value of Y. Normalization removes redundancy by decomposing tables. First Normal Form (1NF) requires atomic values: no repeating groups or multi-valued cells. Second Normal Form (2NF) requires 1NF and that no non-key attribute depends on only part of a composite primary key (no partial dependency)."),
    ("Unit 2: Third Normal Form (3NF) and BCNF", "Third Normal Form (3NF) requires 2NF and no transitive dependency: a non-key attribute must not depend on another non-key attribute. Example: Student(StudentID, Name, DeptID, DeptName). Here StudentID -> DeptID and DeptID -> DeptName, so DeptName depends transitively on the key. Decompose into Student(StudentID, Name, DeptID) and Department(DeptID, DeptName). BCNF is stricter: for every functional dependency X -> Y, X must be a superkey."),
    ("Unit 3: Transactions and ACID", "A transaction is a logical unit of work made of one or more SQL operations. ACID properties: Atomicity means all operations happen or none do. Consistency means a transaction moves the database from one valid state to another. Isolation means concurrent transactions do not interfere with each other. Durability means committed changes survive a crash. Commit makes changes permanent; rollback undoes them."),
    ("Unit 4: Indexing and B+ Trees", "An index is a data structure that speeds up lookups on a column, like the index of a book. A B+ tree keeps keys sorted in a balanced tree whose leaf nodes are linked, which makes range queries efficient. Search, insert and delete take O(log n) time. The trade-off is that indexes use extra storage and slow down writes."),
]
styles = getSampleStyleSheet()
story = []
for i, (title, body) in enumerate(PAGES):
    story += [Paragraph(title, styles["Heading1"]), Spacer(1, 12), Paragraph(body, styles["BodyText"])]
    if i < len(PAGES) - 1:
        story.append(PageBreak())
SimpleDocTemplate(str(out / "dbms_notes.pdf"), pagesize=A4).build(story)

prs = Presentation()
SLIDES = [
    ("SQL Basics: SELECT", "SELECT retrieves rows from a table. Use WHERE to filter rows, ORDER BY to sort and LIMIT to restrict the count. Example: SELECT name FROM students WHERE marks > 80 ORDER BY name;"),
    ("SQL Joins", "A JOIN combines rows from two tables using a related column. INNER JOIN returns only matching rows. LEFT JOIN returns all rows from the left table plus matches from the right. Example: SELECT s.name, d.dept_name FROM students s JOIN departments d ON s.dept_id = d.id;"),
    ("SQL Aggregates", "Aggregate functions summarise many rows into one value: COUNT, SUM, AVG, MIN and MAX. GROUP BY groups rows before aggregating and HAVING filters the groups."),
]
for title, body in SLIDES:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
prs.save(str(out / "sql_basics.pptx"))
print("samples written")
